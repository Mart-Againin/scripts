"""
import_from_existing_reports.py — загружает avg_reach/err/vrpost в
history_db.json НАПРЯМУЮ из уже готовых Excel-отчётов
(output/monthly/monthly_YYYY-MM.xlsx), без обращения к Telegram и без
пересборки чего-либо — просто читает лист "Сводка" каждого файла.

Полезно, если у вас уже накопились месячные отчёты за прошлые месяцы и
не хочется дёргать Telegram заново, чтобы просто перенести уже
посчитанные цифры в базу для Dashboard.

ЧТО БЕРЁТ ИЗ КАЖДОГО ФАЙЛА (лист "Сводка", по каждому каналу):
  - avg_reach = "Охват постов" / "Постов" (в Excel хранится сумма, не
    среднее — здесь пересчитывается в среднее на пост)
  - err       = "Ср. ERR (%)"
  - vrpost    = "Ср. VRpost (%)"

ПРО ПОДПИСЧИКОВ — ОТДЕЛЬНО И ВАЖНО:
  В листе "Сводка" тоже есть колонка "Подписчики", но это число было
  зафиксировано в момент, когда ТОТ отчёт генерировался — то есть оно
  может быть неточным для "конец месяца" (та же проблема, что мы
  разбирали раньше). Поэтому этот скрипт НЕ пишет подписчиков в базу
  автоматически — только выводит их на экран, чтобы вы могли решить,
  использовать ли их через import_subscribers.py (или ввести более
  точные цифры, если они у вас есть из другого источника).

КАК ПОЛЬЗОВАТЬСЯ:
    python import_from_existing_reports.py
        — обработает все monthly_*.xlsx в output/monthly/

    python import_from_existing_reports.py output/monthly/monthly_2026-03.xlsx
        — можно указать конкретные файлы

Как и остальные импорты — сначала показывает список изменений, потом
ОДНО общее подтверждение.
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

from history_db import load_db, save_db, ensure_seeded

MONTHLY_DIR = Path("output/monthly")


def parse_report(path: Path):
    """Возвращает (ym, [(channel, avg_reach, err, vrpost, subscribers), ...])."""
    ym = path.stem.replace("monthly_", "")  # monthly_2026-03.xlsx -> 2026-03
    wb = load_workbook(path, data_only=True)
    if "Сводка" not in wb.sheetnames:
        print(f"  ⚠️  {path.name}: нет листа 'Сводка', пропуск файла")
        return ym, []

    ws = wb["Сводка"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))

    # Заголовок сводки: строка, где колонка B (индекс 1) == "Подписчики"
    header_row_idx = None
    for i, row in enumerate(rows):
        if len(row) > 1 and row[1] == "Подписчики":
            header_row_idx = i
            break
    if header_row_idx is None:
        print(f"  ⚠️  {path.name}: не нашёл строку заголовков в 'Сводка', пропуск")
        return ym, []

    result = []
    for row in rows[header_row_idx + 1:]:
        channel = row[0]
        if not channel or not str(channel).startswith("@"):
            break  # дошли до "ИТОГО / СРЕДНЕЕ" или пустой строки — конец таблицы
        subscribers = row[1]
        posts_count = row[3]
        views_posts = row[5]
        err_frac    = row[13]   # хранится как ДОЛЯ (0.0833), т.к. ячейка
        vrpost_frac = row[15]   # отформатирована как "0.00%" в Excel —
                                 # openpyxl отдаёт сырое число, не "8.33%"

        avg_reach = None
        if isinstance(views_posts, (int, float)) and isinstance(posts_count, (int, float)) and posts_count:
            avg_reach = round(views_posts / posts_count, 1)

        err_pct    = round(err_frac * 100, 2)    if isinstance(err_frac, (int, float)) else None
        vrpost_pct = round(vrpost_frac * 100, 2) if isinstance(vrpost_frac, (int, float)) else None

        result.append((channel, avg_reach, err_pct, vrpost_pct, subscribers))

    return ym, result


def main(paths):
    ensure_seeded()
    db = load_db()

    planned = []  # (ym, ch, new_fields, existing)
    subs_seen = []  # (ym, ch, subscribers_из_файла) — только для информации

    for path in paths:
        ym, rows = parse_report(path)
        if not rows:
            continue
        for ch, avg_reach, err_pct, vrpost_pct, subs in rows:
            new_fields = {}
            if avg_reach is not None:
                new_fields["avg_reach"] = avg_reach
            if isinstance(err_pct, (int, float)):
                new_fields["err"] = round(err_pct, 2)
            if isinstance(vrpost_pct, (int, float)):
                new_fields["vrpost"] = round(vrpost_pct, 2)
            if not new_fields:
                continue
            existing = db.get(ym, {}).get(ch)
            planned.append((ym, ch, new_fields, existing))
            if isinstance(subs, (int, float)):
                subs_seen.append((ym, ch, int(subs)))

    if not planned:
        print("Нечего записывать — не нашёл подходящих данных ни в одном файле.")
        return

    print(f"\nБудет применено изменений (avg_reach/err/vrpost): {len(planned)}\n")
    for ym, ch, new_fields, existing in planned:
        if existing:
            old_vals = {k: existing.get(k) for k in new_fields}
            print(f"  {ch} {ym}: {new_fields}  (было: {old_vals})")
        else:
            print(f"  {ch} {ym}: новая запись {new_fields}")

    answer = input(f"\nПрименить все {len(planned)} изменений? (да/нет): ").strip().lower()
    if answer in ("да", "yes", "y", "д"):
        for ym, ch, new_fields, existing in planned:
            db.setdefault(ym, {})
            if existing:
                existing.update(new_fields)
            else:
                db[ym][ch] = new_fields
        save_db(db)
        print(f"\nГотово — записано изменений: {len(planned)}.")
    else:
        print("Отменено, ничего не записано.")

    if subs_seen:
        print("\n--- Подписчики, найденные в файлах (НЕ записаны, см. предупреждение выше) ---")
        for ym, ch, subs in subs_seen:
            print(f"  {ym},{ch},{subs}")
        print("\nЕсли хотите их использовать — скопируйте нужные строки в subscribers.csv "
              "(формат: month,channel,subscribers) и запустите import_subscribers.py.")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        files = [Path(p) for p in sys.argv[1:]]
    else:
        files = sorted(MONTHLY_DIR.glob("monthly_*.xlsx"))
    if not files:
        print(f"Не нашёл ни одного monthly_*.xlsx в {MONTHLY_DIR}")
        sys.exit(1)
    main(files)
